"""PowerPoint -> markdown: one `## Slide N` section per slide, speaker notes kept."""
from __future__ import annotations

import io

from pal.input_converters.components.base import ConversionError, table_to_markdown


def _shape_lines(shape, out: list[str]) -> None:
    if getattr(shape, "shape_type", None) == 6 or hasattr(shape, "shapes"):  # group
        for s in getattr(shape, "shapes", []):
            _shape_lines(s, out)
        return
    if getattr(shape, "has_table", False) and shape.has_table:
        rows = [[c.text for c in r.cells] for r in shape.table.rows]
        md = table_to_markdown(rows)
        if md:
            out.append(md)
        return
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            t = "".join(r.text for r in para.runs).strip() or para.text.strip()
            if t:
                out.append(("  " * min(para.level, 6)) + "- " + t if para.level else t)


def convert(raw: bytes) -> str:
    try:
        from pptx import Presentation
    except ImportError:
        raise ConversionError("PowerPoint support needs python-pptx (pip install python-pptx)") from None
    try:
        prs = Presentation(io.BytesIO(raw))
    except Exception as e:
        raise ConversionError(f"could not read PPTX ({e})") from e
    sections = []
    for i, slide in enumerate(prs.slides, 1):
        title = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text_frame.text.strip()
        lines: list[str] = []
        for shape in slide.shapes:
            if slide.shapes.title is not None and shape.shape_id == slide.shapes.title.shape_id:
                continue
            _shape_lines(shape, lines)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"**Notes:** {notes}")
        head = f"## Slide {i}" + (f": {title}" if title else "")
        sections.append(head + ("\n\n" + "\n\n".join(lines) if lines else ""))
    return "\n\n".join(sections)
