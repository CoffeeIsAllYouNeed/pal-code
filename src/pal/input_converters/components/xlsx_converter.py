"""Excel -> markdown: one `## Sheet: name` section with a pipe table per sheet.
Formula cells give their cached (last-saved) value."""
from __future__ import annotations

import io

from pal.input_converters.components.base import ConversionError, table_to_markdown


def _fmt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if hasattr(v, "isoformat"):
        return v.isoformat(sep=" ") if hasattr(v, "hour") and (v.hour or v.minute or v.second) else v.isoformat()[:10]
    return str(v)


def convert(raw: bytes) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ConversionError("Excel support needs openpyxl (pip install openpyxl)") from None
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as e:
        raise ConversionError(f"could not read XLSX ({e})") from e
    sections = []
    try:
        for ws in wb.worksheets:
            rows = [[_fmt(v) for v in row] for row in ws.iter_rows(values_only=True)]
            table = table_to_markdown(rows)
            if table:
                sections.append(f"## Sheet: {ws.title}\n\n{table}")
    finally:
        wb.close()
    return "\n\n".join(sections)
