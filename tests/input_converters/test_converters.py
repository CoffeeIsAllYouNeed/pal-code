import io
import shutil
import subprocess

import pytest
from tests.helpers import client, make_service

from pal.input_converters import ConversionError, convert_to_markdown

needs_pandoc = pytest.mark.skipif(shutil.which("pandoc") is None, reason="pandoc not installed")


def pandoc_make(fmt: str, source: str) -> bytes:
    """Build a fixture in `fmt` (docx/odt/epub/rtf) from markdown using pandoc."""
    return subprocess.run(["pandoc", "-f", "markdown", "-t", fmt, "-o", "-"], input=source.encode(),
                          capture_output=True, check=True).stdout


MD = "# Title\n\nHello **world**.\n\n- one\n- two\n\n| a | b |\n|---|---|\n| 1 | 2 |\n"


@needs_pandoc
def test_html_becomes_markdown_without_scripts_or_tags():
    html = (b"<html><head><style>p{color:red}</style><script>alert(1)</script></head>"
            b"<body><h1>Title</h1><p>Hello <b>world</b></p><ul><li>one</li></ul></body></html>")
    r = convert_to_markdown("page.html", html)
    assert r.converted and r.source_format == "html"
    assert "# Title" in r.text and "**world**" in r.text and "- one" in r.text
    assert "<" not in r.text and "alert" not in r.text and "color:red" not in r.text


@needs_pandoc
@pytest.mark.parametrize("ext", ["docx", "odt", "epub", "rtf"])
def test_pandoc_document_formats(ext):
    r = convert_to_markdown(f"doc.{ext}", pandoc_make(ext, MD))
    assert r.converted and "Hello" in r.text and "world" in r.text and "one" in r.text


@needs_pandoc
def test_docx_keeps_headings_and_tables():
    r = convert_to_markdown("d.docx", pandoc_make("docx", MD))
    assert "# Title" in r.text and "| a" in r.text and "| 1" in r.text


@needs_pandoc
def test_latex():
    tex = rb"\documentclass{article}\begin{document}\section{Intro}Some \textbf{bold} text.\end{document}"
    r = convert_to_markdown("paper.tex", tex)
    assert "Intro" in r.text and "**bold**" in r.text and "\\textbf" not in r.text


def make_pdf(text: str) -> bytes:
    # minimal one-page PDF with a text layer
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
            b"/Resources << /Font << /F1 5 0 R >> >> >>",
            b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offs = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        out += b"%d 0 obj\n" % i + o + b"\nendobj\n"
    x = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offs)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF" % (len(objs) + 1, x)
    return out


def test_pdf_with_text_layer():
    r = convert_to_markdown("a.pdf", make_pdf("Hello PDF world"))
    assert r.converted and "Hello PDF world" in r.text


def test_pdf_without_text_layer_is_rejected():
    from pypdf import PdfWriter
    w = PdfWriter()
    w.add_blank_page(200, 200)
    buf = io.BytesIO()
    w.write(buf)
    with pytest.raises(ConversionError, match="no text layer"):
        convert_to_markdown("scan.pdf", buf.getvalue())


def test_pptx_slides_notes_and_tables():
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Roadmap"
    s.placeholders[1].text_frame.text = "Ship converters"
    s.notes_slide.notes_text_frame.text = "say this aloud"
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Numbers"
    t = s2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    for (i, j), v in {(0, 0): "k", (0, 1): "v", (1, 0): "x", (1, 1): "42"}.items():
        t.cell(i, j).text = v
    buf = io.BytesIO()
    prs.save(buf)
    r = convert_to_markdown("deck.pptx", buf.getvalue())
    assert "## Slide 1: Roadmap" in r.text and "Ship converters" in r.text and "say this aloud" in r.text
    assert "## Slide 2: Numbers" in r.text and "| x | 42 |" in r.text


def test_xlsx_sheets_become_tables():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["item", "qty"])
    ws.append(["a|b", 3.0])
    wb.create_sheet("Empty")
    wb.create_sheet("More").append(["x"])
    buf = io.BytesIO()
    wb.save(buf)
    r = convert_to_markdown("d.xlsx", buf.getvalue())
    assert "## Sheet: Sales" in r.text and "| item | qty |" in r.text
    assert "a\\|b" in r.text and "| 3 |" in r.text
    assert "Empty" not in r.text and "## Sheet: More" in r.text


def test_plain_text_passes_through_unchanged():
    r = convert_to_markdown("notes.py", b"print('hi')\n<b>keep</b>\n")
    assert not r.converted and r.source_format == "text" and "<b>keep</b>" in r.text


def test_unknown_binary_and_legacy_are_rejected_with_reason():
    with pytest.raises(ConversionError, match="binary files are not supported"):
        convert_to_markdown("x.bin", b"\x00\x01\x02")
    with pytest.raises(ConversionError, match="save it as .docx"):
        convert_to_markdown("old.doc", b"\xd0\xcf\x11\xe0")


def test_corrupt_file_gives_conversion_error_not_crash():
    with pytest.raises(ConversionError):
        convert_to_markdown("bad.xlsx", b"not a zip")
    with pytest.raises(ConversionError):
        convert_to_markdown("bad.pptx", b"not a zip")
    with pytest.raises(ConversionError):
        convert_to_markdown("bad.pdf", b"not a pdf")


def test_empty_file_rejected():
    with pytest.raises(ConversionError, match="no readable text"):
        convert_to_markdown("e.txt", b"   \n")


def test_flask_upload_converts_docx_before_api(tmp_path):
    if shutil.which("pandoc") is None:
        pytest.skip("pandoc not installed")
    svc, api = make_service(tmp_path)
    r = client(svc).post("/api/chat", data={
        "message": "read the file", "provider": "anthropic", "model": "m", "api_key": "k", "mode": "steps",
        "methods": ["remove_articles"], "files": (io.BytesIO(pandoc_make("docx", MD)), "brief.docx")})
    d = r.get_json()
    assert r.status_code == 200 and not d["file_errors"]
    assert "[file: brief.docx | converted from docx to markdown]" in d["sent"] and "# Title" in d["sent"] and "PK" not in d["sent"]


def test_conversion_errors_are_user_errors_with_a_400():
    from pal.errors import UserInputError
    assert issubclass(ConversionError, UserInputError) and ConversionError("x").http_status == 400
